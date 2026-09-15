#!/usr/bin/env python3
"""
Bundle Markdown files into one .docx for reading away from a terminal.

Written because the legal pack was exported once by hand and then needed
exporting again the moment anything changed — and an export nobody can repeat
goes stale silently, which for a document someone is about to rely on is worse
than not having it.

    tools/md-to-docx.py OUT.docx "Title" FILE.md [FILE.md ...]

Handles the subset this project's documents actually use: ATX headings, bullet
and numbered lists, fenced code, pipe tables, block quotes, horizontal rules,
and inline `code`, **bold** and *italic*. Anything else is emitted as plain
text rather than dropped — a converter that silently loses a clause is not one
you can hand to a lawyer.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

INLINE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`|~~.+?~~)")


def add_runs(par, text):
    """Inline formatting. Unmatched markers stay as literal text on purpose."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("~~") and part.endswith("~~") and len(part) > 4:
            r = par.add_run(part[2:-2]); r.font.strike = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x8B, 0x22, 0x52)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def is_table_row(line):
    return line.strip().startswith("|") and line.strip().endswith("|")


def cells(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def convert(doc, path):
    lines = Path(path).read_text().split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # fenced code
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run("\n".join(buf))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            continue

        # pipe table
        if is_table_row(line) and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|-: "):
            header = cells(line)
            i += 2
            rows = []
            while i < len(lines) and is_table_row(lines[i]):
                rows.append(cells(lines[i])); i += 1
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Light Grid Accent 1"
            for c, text in zip(t.rows[0].cells, header):
                c.text = ""
                add_runs(c.paragraphs[0], text)
                for r in c.paragraphs[0].runs:
                    r.bold = True
            for row in rows:
                rc = t.add_row().cells
                for c, text in zip(rc, row + [""] * (len(header) - len(row))):
                    c.text = ""
                    add_runs(c.paragraphs[0], text)
            doc.add_paragraph()
            continue

        if not stripped:
            i += 1
            continue

        if stripped in ("---", "***", "___"):
            doc.add_paragraph().add_run("─" * 40).font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            doc.add_heading(m.group(2).strip(), min(len(m.group(1)), 4))
            i += 1
            continue

        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(24)
            add_runs(p, stripped.lstrip("> ").strip())
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        m = re.match(r"^[-*+]\s+(.*)$", stripped)
        if m:
            add_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
            i += 1
            continue

        m = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if m:
            add_runs(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue

        # paragraph: join continuation lines
        buf = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^(#{1,6}\s|[-*+]\s|\d+[.)]\s|>|```|\|)", lines[i].strip()
        ) and lines[i].strip() not in ("---", "***", "___"):
            buf.append(lines[i].strip()); i += 1
        add_runs(doc.add_paragraph(), " ".join(buf))


def main(argv):
    if len(argv) < 4:
        print(__doc__)
        return 2
    out, title, files = argv[1], argv[2], argv[3:]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(title)
    r.bold = True
    r.font.size = Pt(22)

    from datetime import date
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = s.add_run(f"Generated {date.today().isoformat()} from the repository — "
                   f"regenerate with tools/md-to-docx.py rather than editing this file")
    sr.italic = True
    sr.font.size = Pt(9)
    sr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    for n, f in enumerate(files):
        doc.add_page_break()
        convert(doc, f)
    doc.save(out)
    print(f"{out}  ({len(files)} documents)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
